import json
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# 1. Load Categories
with open('moder-store.categories.updated.json', 'r', encoding='utf-8') as f:
    cat_list = json.load(f)

categories_map = {}
for c in cat_list:
    cid = c.get('_id')
    if isinstance(cid, dict):
        cid = cid.get('$oid')
    if cid:
        categories_map[str(cid)] = c.get('name', 'Uncategorized')

# 2. Load Inventories (from the user's downloaded database export)
with open('moder-store.inventories.json', 'r', encoding='utf-8') as f:
    inventories = json.load(f)

inventory_map = {}
for inv in inventories:
    pid = inv.get('productId')
    if isinstance(pid, dict):
        pid = pid.get('$oid')
    if pid:
        inventory_map[str(pid)] = {
            'quantityInStock': inv.get('quantityInStock', 0),
            'soldQuantity': inv.get('soldQuantity', 0),
            'totalStock': inv.get('totalStock', 0),
            'unit': inv.get('unit', ''),
            'sku': inv.get('SKU', '')
        }

print(f"Loaded {len(inventory_map)} inventory records from moder-store.inventories.json")

# 3. Load 480 Products
with open('moder-store.products.updated.json', 'r', encoding='utf-8') as f:
    products_raw = json.load(f)

print(f"Loaded {len(products_raw)} products from moder-store.products.updated.json")

# 4. Build unified 480 products dataset with real stock
processed_products = []

for idx, p in enumerate(products_raw, start=1):
    pid = p.get('_id')
    if isinstance(pid, dict):
        pid = pid.get('$oid')
    pid_str = str(pid)

    name = p.get('name', 'Unnamed Product')
    sub_name = p.get('subName', '')
    sku = p.get('sku') or f"SKU-{pid_str[-6:].upper()}"
    unit = p.get('unit', 'Unit')
    base_price = float(p.get('basePrice', 0))
    discount = float(p.get('discountPercentage', 0))
    selling_price = round(base_price * (1 - discount / 100), 2) if discount > 0 else base_price

    cat_id = p.get('category')
    if isinstance(cat_id, dict):
        cat_id = cat_id.get('$oid')
    cat_name = categories_map.get(str(cat_id), 'General')

    # Check real stock
    inv_info = inventory_map.get(pid_str)
    if inv_info:
        current_stock = inv_info['quantityInStock']
        sold_qty = inv_info['soldQuantity']
        total_stock = inv_info['totalStock']
        has_stock_record = True
    else:
        current_stock = 0
        sold_qty = 0
        total_stock = 0
        has_stock_record = False

    processed_products.append({
        'slNo': idx,
        'id': pid_str,
        'name': name,
        'subName': sub_name,
        'sku': sku,
        'category': cat_name,
        'unit': unit,
        'basePrice': base_price,
        'sellingPrice': selling_price,
        'discountPercentage': discount,
        'currentStock': current_stock,
        'soldQuantity': sold_qty,
        'totalStock': total_stock,
        'stockToAdd': 0,
        'hasStockRecord': has_stock_record
    })

# Save unified JSON
with open('all_480_products_stock.json', 'w', encoding='utf-8') as f:
    json.dump(processed_products, f, indent=2, ensure_ascii=False)

print(f"Successfully generated all_480_products_stock.json with {len(processed_products)} items.")
print(f"Products with stock > 0 in DB: {sum(1 for p in processed_products if p['currentStock'] > 0)}")
print(f"Products with 0 stock in DB: {sum(1 for p in processed_products if p['currentStock'] == 0)}")

# 5. Create Professional Excel with Real Data
wb = openpyxl.Workbook()
ws_stock = wb.active
ws_stock.title = "Stock Update Sheet"

# Styling definitions
header_fill = PatternFill(start_color="1B5E20", end_color="1B5E20", fill_type="solid")
header_fill_accent = PatternFill(start_color="E65100", end_color="E65100", fill_type="solid")
header_fill_formula = PatternFill(start_color="0D47A1", end_color="0D47A1", fill_type="solid")
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

zebra_fill = PatternFill(start_color="F9FBF9", end_color="F9FBF9", fill_type="solid")
input_fill = PatternFill(start_color="FFF8E1", end_color="FFF8E1", fill_type="solid")

thin_border = Border(
    left=Side(style='thin', color='E0E0E0'),
    right=Side(style='thin', color='E0E0E0'),
    top=Side(style='thin', color='E0E0E0'),
    bottom=Side(style='thin', color='E0E0E0')
)

headers = [
    "Sl No",
    "Product ID (Do Not Change)",
    "Product Name",
    "Malayalam / Sub Name",
    "SKU Code",
    "Category",
    "Unit",
    "Base Price (₹)",
    "Selling Price (₹)",
    "Current Stock in DB",
    "Sold Qty",
    "Stock To Add (Type New Qty Here)",
    "Projected Total Stock (=Current+Add)",
    "Status / Remarks"
]

ws_stock.append(headers)

for col_idx in range(1, len(headers) + 1):
    cell = ws_stock.cell(row=1, column=col_idx)
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if col_idx == 12: # Stock To Add
        cell.fill = header_fill_accent
    elif col_idx == 13: # Projected Total
        cell.fill = header_fill_formula
    else:
        cell.fill = header_fill

ws_stock.row_dimensions[1].height = 32

for p in processed_products:
    row_num = p['slNo'] + 1
    # Column 10: Current Stock, Column 12: Stock To Add
    formula_str = f"=J{row_num}+L{row_num}"
    
    if p['currentStock'] == 0:
        status = "No Stock Added Yet"
    elif p['currentStock'] <= 5:
        status = "Low Stock"
    else:
        status = "In Stock"

    row_data = [
        p['slNo'],
        p['id'],
        p['name'],
        p['subName'],
        p['sku'],
        p['category'],
        p['unit'],
        p['basePrice'],
        p['sellingPrice'],
        p['currentStock'],
        p['soldQuantity'],
        0, # Stock To Add
        formula_str,
        status
    ]
    ws_stock.append(row_data)

    for col_idx in range(1, len(row_data) + 1):
        cell = ws_stock.cell(row=row_num, column=col_idx)
        cell.border = thin_border
        
        # Center aligns
        if col_idx in [1, 5, 7, 14]:
            cell.alignment = Alignment(horizontal="center", vertical="center")
        elif col_idx in [8, 9, 10, 11, 12, 13]:
            cell.alignment = Alignment(horizontal="right", vertical="center")
        else:
            cell.alignment = Alignment(horizontal="left", vertical="center")

        # Color coding
        if col_idx == 12:
            cell.fill = input_fill
            cell.font = Font(name="Calibri", size=11, bold=True, color="B78103")
        elif col_idx == 13:
            cell.font = Font(name="Calibri", size=11, bold=True, color="0D47A1")
        elif row_num % 2 == 1:
            cell.fill = zebra_fill

    ws_stock.row_dimensions[row_num].height = 22

# Column widths
column_widths = {
    1: 8,   # Sl No
    2: 28,  # Product ID
    3: 32,  # Product Name
    4: 26,  # Sub Name
    5: 16,  # SKU
    6: 22,  # Category
    7: 12,  # Unit
    8: 14,  # Base Price
    9: 16,  # Selling Price
    10: 20, # Current Stock in DB
    11: 14, # Sold Qty
    12: 28, # Stock To Add
    13: 26, # Projected Total
    14: 22  # Status
}

for col_idx, width in column_widths.items():
    col_letter = get_column_letter(col_idx)
    ws_stock.column_dimensions[col_letter].width = width

ws_stock.freeze_panes = "C2"

# 6. Instructions Sheet
ws_guide = wb.create_sheet(title="Instructions & Guide")
guide_data = [
    ["MODERN STORE - 480 PRODUCTS BULK STOCK UPDATE INSTRUCTIONS"],
    [""],
    ["Column Header", "Description", "Required / Action", "Example"],
    ["Product ID (Do Not Change)", "Exact MongoDB ObjectID required by backend API.", "LOCKED (DO NOT EDIT)", "6992a174f62501122c4c8c0c"],
    ["Product Name", "Real registered product title from database.", "Read Only", "Urulakizhangu"],
    ["Malayalam / Sub Name", "Malayalam / Regional search name.", "Read Only", "potato, ഉരുളക്കിഴങ്"],
    ["SKU Code", "Unique Stock Keeping Unit.", "Read Only", "VEG-POT-1198"],
    ["Category", "Store Category (Vegetables, Fruits, etc.).", "Read Only", "Vegetable"],
    ["Unit", "Unit of measurement (KG, PACK, PCS, etc.).", "Read Only", "KG"],
    ["Current Stock in DB", "Real stock remaining as per moder-store.inventories.json.", "Database Value (Read Only)", "11 (or 0 if not added yet)"],
    ["Stock To Add (Type New Qty Here)", "ENTER THE NEW STOCK QUANTITY HERE!", "TYPE NUMBER HERE", "Enter 50 to add 50 units"],
    ["Projected Total Stock", "Automatic formula: =Current Stock + Stock To Add.", "Formula (=J2+L2)", "Auto-calculates"],
    ["Status / Remarks", "In Stock / Low Stock / No Stock Added Yet.", "Information", "No Stock Added Yet"]
]

for r_idx, row in enumerate(guide_data, 1):
    ws_guide.append(row)
    for c_idx in range(1, len(row) + 1):
        cell = ws_guide.cell(row=r_idx, column=c_idx)
        if r_idx == 1:
            cell.font = Font(name="Calibri", size=14, bold=True, color="1B5E20")
        elif r_idx == 3:
            cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")

guide_widths = {1: 30, 2: 55, 3: 25, 4: 30}
for col_idx, width in guide_widths.items():
    ws_guide.column_dimensions[get_column_letter(col_idx)].width = width

excel_path = "ModernStore_All_480_Products_Stock.xlsx"
wb.save(excel_path)
print(f"Saved real Excel file to {excel_path}")
