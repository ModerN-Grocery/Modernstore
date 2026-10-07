import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import json
import random

# Categories & Real Catalog Data
catalog = {
    "Fresh Fruits": [
        ("Shimla Royal Apple", "KG", 180, 150),
        ("Robusta Banana", "KG", 60, 48),
        ("Alphonsa Mango Ratnagiri", "KG", 240, 210),
        ("Nagpur Sweet Orange", "KG", 90, 75),
        ("Seedless Green Grapes", "KG", 130, 110),
        ("Black Seedless Grapes", "KG", 145, 125),
        ("Fresh Kiran Watermelon", "KG", 40, 32),
        ("Red Lady Papaya", "KG", 55, 45),
        ("Pomegranate Kesar", "KG", 220, 185),
        ("Pineapple Queen", "PCS", 85, 70),
        ("Kashmiri Red Delicious Apple", "KG", 210, 180),
        ("Yellaki Small Banana", "KG", 90, 75),
        ("Mosambi Sweet Lime", "KG", 80, 68),
        ("Imported Zespri Kiwi (Pack of 3)", "BOX", 130, 105),
        ("Dragon Fruit Red", "PCS", 110, 89),
        ("Fresh Allahabad Guava", "KG", 70, 58),
        ("Custard Apple (Seethaphal)", "KG", 140, 115),
        ("Raw Green Mango (Totapuri)", "KG", 60, 50),
        ("Sapota Brown (Chikoo)", "KG", 65, 52),
        ("Hass Avocado Premium", "PCS", 95, 79),
        ("Strawberry Mahabaleshwar 200g", "BOX", 95, 82),
        ("Fresh Blueberries 125g", "BOX", 260, 225),
        ("Sweet Melon (Chibudh)", "KG", 60, 48),
        ("Plums Red Imported", "KG", 280, 240),
    ],
    "Fresh Vegetables": [
        ("Fresh Tomato Local Country", "KG", 40, 32),
        ("Fresh Tomato Hybrid", "KG", 38, 30),
        ("Red Onion Big Nasik", "KG", 55, 45),
        ("Sambar Small Onion (Shallots)", "KG", 90, 78),
        ("Potato Jyoti", "KG", 35, 28),
        ("Carrot Ooty Fresh", "KG", 70, 58),
        ("Green Beans Fine", "KG", 85, 72),
        ("Ladies Finger Tender (Bhindi)", "KG", 60, 48),
        ("Egg Plant (Brinjal Black Big)", "KG", 45, 36),
        ("Brinjal Varikatri Small", "KG", 50, 40),
        ("Green Cabbage Fresh", "KG", 35, 28),
        ("Cauliflower Snowball", "PCS", 50, 42),
        ("Garlic Desi Small Pods", "KG", 240, 210),
        ("Green Chilli Spicy Hot", "KG", 80, 68),
        ("Ginger Local Fresh", "KG", 160, 135),
        ("Curry Leaves Fresh Bunch", "PACKET", 15, 10),
        ("Coriander Leaves Bunch", "PACKET", 20, 15),
        ("Mint Leaves (Pudina)", "PACKET", 20, 15),
        ("Palak (Spinach Fresh)", "PACKET", 30, 24),
        ("Bottle Gourd (Lauki)", "KG", 40, 32),
        ("Bitter Gourd Small (Karela)", "KG", 65, 52),
        ("Snake Gourd (Padavalanga)", "KG", 45, 38),
        ("Cucumber Green Salad", "KG", 40, 30),
        ("Capsicum Green Bell Pepper", "KG", 90, 75),
        ("Capsicum Red/Yellow Bell", "KG", 220, 185),
        ("Beetroot Deep Red", "KG", 50, 42),
        ("Radish White (Mullangi)", "KG", 40, 32),
        ("Drumsticks Long (Muringakka)", "KG", 120, 98),
        ("Raw Banana (Nendran Green)", "KG", 75, 62),
        ("Elephant Foot Yam (Chena)", "KG", 60, 50),
        ("Colocasia Taro (Chembu)", "KG", 70, 58),
        ("Ash Gourd (Kumbalanga)", "KG", 35, 28),
        ("Pumpkin Yellow (Mathanga)", "KG", 30, 24),
        ("Mushroom Button Pack 200g", "PACKET", 60, 50),
        ("Sweet Corn Whole Cob", "PCS", 30, 25),
        ("Lemon Yellow Juicy (6 Pcs)", "PACKET", 40, 32),
    ],
    "Dairy & Breakfast": [
        ("Milma Toned Milk 500ml", "PACKET", 27, 27),
        ("Milma Rich Homogenized Milk 500ml", "PACKET", 30, 30),
        ("Amul Taaza Milk 1L Tetra", "PACKET", 74, 70),
        ("Amul Salted Butter 100g", "PCS", 60, 58),
        ("Amul Salted Butter 500g", "PCS", 285, 275),
        ("Amul Fresh Malai Paneer 200g", "PACKET", 95, 88),
        ("Milma Fresh Curd Pouch 500g", "PACKET", 35, 33),
        ("White Table Eggs (Pack of 6)", "BOX", 54, 48),
        ("White Table Eggs (Pack of 12)", "BOX", 105, 92),
        ("White Eggs Farm Fresh (Crate of 30)", "BOX", 210, 190),
        ("Brown Country Eggs Organic (Pack of 6)", "BOX", 85, 75),
        ("Amul Processed Cheese Slices 200g", "PACKET", 150, 138),
        ("Amul Processed Cheese Block 200g", "PCS", 140, 128),
        ("Milma Pure Agmark Ghee 500ml", "PCS", 350, 335),
        ("Amul Fresh Cream 250ml", "PACKET", 68, 62),
        ("Nestle Milkmaid Sweetened 400g", "PCS", 148, 140),
        ("Epigamia Greek Yogurt Strawberry 100g", "PCS", 55, 48),
        ("ID Fresh Idly & Dosa Batter 1kg", "PACKET", 95, 82),
        ("Modern White Sandwich Bread 400g", "PACKET", 45, 42),
        ("Modern 100% Whole Wheat Bread 400g", "PACKET", 55, 50),
        ("Britannia Little Hearts Biscuits 75g", "PACKET", 25, 22),
    ],
    "Beverages & Cool Drinks": [
        ("Sprite Lime 750ml Pet", "PCS", 45, 40),
        ("Sprite Lime 2L Family Bottle", "PCS", 95, 85),
        ("Coca Cola Original 750ml", "PCS", 45, 40),
        ("Coca Cola 2L Pet Bottle", "PCS", 95, 85),
        ("Thums Up Charged 750ml", "PCS", 45, 40),
        ("Pepsi 750ml Pet", "PCS", 40, 38),
        ("Mirinda Orange 750ml", "PCS", 40, 38),
        ("7 Up Lemonade 750ml", "PCS", 40, 38),
        ("Real Fruit Power Mango 1L", "PCS", 130, 115),
        ("Real Fruit Power Mixed Fruit 1L", "PCS", 130, 115),
        ("Real Fruit Power Guava 1L", "PCS", 130, 115),
        ("Tropicana 100% Orange Juice 1L", "PCS", 135, 120),
        ("Frooti Fresh 'N' Juicy Mango 150ml", "PACKET", 15, 14),
        ("Red Bull Energy Drink Can 250ml", "PCS", 125, 120),
        ("Kinley Purified Water 1L", "PCS", 20, 20),
        ("Bisleri Mineral Water 2L", "PCS", 30, 30),
        ("Appy Fizz Sparkling Apple Drink 600ml", "PCS", 40, 36),
        ("Mountain Dew Soft Drink 750ml", "PCS", 45, 40),
        ("Bovonto Soft Drink 750ml", "PCS", 40, 38),
    ],
    "Staples, Rice & Atta": [
        ("Palakkadan Matta Rice Long Grain 5kg", "PACKET", 330, 298),
        ("Palakkadan Matta Rice Long Grain 10kg", "PACKET", 640, 580),
        ("Thanjavur Ponni Boiled Rice 5kg", "PACKET", 340, 310),
        ("India Gate Basmati Rice Feast 1kg", "PACKET", 160, 140),
        ("Daawat Rozana Gold Basmati Rice 5kg", "PACKET", 520, 460),
        ("Aashirvaad Sharbati Atta 5kg", "PACKET", 295, 270),
        ("Chakki Fresh Whole Wheat Atta 10kg", "PACKET", 510, 465),
        ("Nirapara Roasted Rice Powder 1kg", "PACKET", 85, 75),
        ("Double Horse Easy Pathiri Powder 1kg", "PACKET", 90, 80),
        ("Toor Dal Premium Unpolished 1kg", "PACKET", 175, 158),
        ("Moong Dal Yellow Split 1kg", "PACKET", 148, 132),
        ("Urad Dal Gota White 1kg", "PACKET", 165, 148),
        ("Chana Dal Bengal Gram 1kg", "PACKET", 112, 99),
        ("Kala Chana Desi 1kg", "PACKET", 118, 105),
        ("Kabuli Chana Dollar Big 1kg", "PACKET", 185, 165),
        ("Fortune Sunlite Refined Sunflower Oil 1L", "PACKET", 142, 130),
        ("KLF Nirmal Pure Coconut Oil 1L Pouch", "PACKET", 245, 225),
        ("Idhayam Pure Sesame Gingelly Oil 500ml", "PCS", 215, 198),
        ("Tata Salt Vacuum Evaporated Iodized 1kg", "PACKET", 28, 26),
        ("Madhur Pure & Hygienic Sugar 1kg", "PACKET", 52, 46),
        ("Solid Round Jaggery (Vellam) 1kg", "PACKET", 92, 80),
        ("Roasted Sooji Rava 1kg", "PACKET", 68, 60),
    ],
    "Snacks & Instant Food": [
        ("Parle-G Gold Glucose Biscuits 1kg", "PACKET", 140, 128),
        ("Britannia Good Day Butter Cookies 600g", "PACKET", 138, 122),
        ("Oreo Vanilla Creme Sandwich 300g", "PACKET", 90, 82),
        ("Britannia Marie Gold 1kg Family Pack", "PACKET", 155, 138),
        ("Lay's India's Magic Masala 50g", "PACKET", 20, 20),
        ("Kurkure Masala Munch 90g", "PACKET", 30, 28),
        ("Authentic Kerala Banana Chips 250g", "PACKET", 145, 125),
        ("Crispy Kerala Tapioca Chips 200g", "PACKET", 85, 72),
        ("Haldiram Nagpur Bhujia Sev 400g", "PACKET", 135, 118),
        ("Haldiram All In One Mixture 400g", "PACKET", 135, 118),
        ("Cadbury Dairy Milk Silk Plain 150g", "PCS", 185, 175),
        ("Nestle KitKat 4 Finger Chocolate 38g", "PCS", 30, 30),
        ("Maggi 2-Minute Masala Noodles 12-Pack", "PACKET", 172, 155),
        ("Knorr Classic Thick Tomato Soup 4-Pack", "PACKET", 68, 60),
        ("Chings Secret Hakka Noodles 600g", "PACKET", 95, 82),
        ("Kissan Fresh Tomato Ketchup 1kg", "PCS", 155, 135),
    ],
    "Spices & Masalas": [
        ("Eastern Special Chilli Powder 500g", "PACKET", 195, 175),
        ("Eastern Coriander Powder 500g", "PACKET", 155, 138),
        ("Eastern Turmeric Powder 250g", "PACKET", 78, 69),
        ("Eastern Malabar Sambar Powder 200g", "PACKET", 84, 76),
        ("Eastern Meat Masala Powder 200g", "PACKET", 90, 82),
        ("Eastern Chicken Masala 200g", "PACKET", 90, 82),
        ("Eastern Fish Curry Masala 200g", "PACKET", 80, 72),
        ("Black Pepper Whole Tellicherry 100g", "PACKET", 125, 110),
        ("Green Cardamom Whole (Elaichi) 50g", "PACKET", 185, 165),
        ("Cloves Whole (Grampoo) 50g", "PACKET", 98, 88),
        ("Cinnamon Sticks Cassia 100g", "PACKET", 85, 75),
        ("Black Mustard Seeds (Kaduku) 250g", "PACKET", 48, 40),
        ("Cumin Seeds (Jeerakam) 200g", "PACKET", 115, 102),
        ("Fenugreek Seeds (Uluva) 200g", "PACKET", 42, 36),
        ("Kashmiri Bright Red Chilli 250g", "PACKET", 165, 148),
        ("Fennel Seeds (Perunjeerakam) 200g", "PACKET", 95, 82),
    ],
    "Personal Care & Cleaning": [
        ("Surf Excel Easy Wash Detergent 1kg", "PACKET", 148, 138),
        ("Ariel Matic Front Load Mat Detergent 2kg", "PACKET", 495, 445),
        ("Vim Dishwash Gel Lemon 750ml", "PCS", 168, 150),
        ("Vim Dishwash Bar 300g Pack of 3", "PACKET", 62, 55),
        ("Dettol Antiseptic Disinfectant Liquid 550ml", "PCS", 228, 212),
        ("Lizol Surface Cleaner Citrus 2L", "PCS", 395, 355),
        ("Harpic Power Plus Toilet Cleaner 1L", "PCS", 218, 198),
        ("Comfort Morning Fresh Fabric Conditioner 860ml", "PCS", 238, 215),
        ("Good Knight Gold Flash Mosquito Machine + Refill", "BOX", 125, 110),
        ("Pril Lime Dishwash Liquid 425ml", "PCS", 115, 102),
        ("Colgate MaxFresh Spicy Fresh Toothpaste 150g", "PCS", 120, 105),
        ("Lifebuoy Total Germ Protection Soap 125g (Pack 4)", "PACKET", 160, 142),
        ("Dettol Original Bathing Soap 125g (Pack 4)", "PACKET", 195, 175),
        ("Whisper Choice Ultra Wings Sanitary Pads (20 Pads)", "PACKET", 165, 148),
        ("Head & Shoulders Anti-Dandruff Shampoo 340ml", "PCS", 310, 275),
    ]
}

# Expand realistically to reach exactly 480 items
base_products = []
for cat, items in catalog.items():
    for name, unit, mrp, sp in items:
        base_products.append({
            'name': name,
            'category': cat,
            'unit': unit,
            'mrp': mrp,
            'sp': sp
        })

final_480 = []
tier_tags = ["", " (Value Saver)", " (Family Pack)", " (Special Edition)", " (Economy Pack)", " (Combo Offer)"]
base_count = len(base_products)

id_hex_base = 0x67ec29000000000000000000
random.seed(101)

for i in range(1, 481):
    src = base_products[(i - 1) % base_count]
    tier_idx = (i - 1) // base_count
    suffix = tier_tags[tier_idx % len(tier_tags)]
    
    prod_name = src['name'] if tier_idx == 0 else f"{src['name']}{suffix}"
    multiplier = 1.0 if tier_idx == 0 else (1.0 + (tier_idx * 0.4))
    mrp = int(src['mrp'] * multiplier)
    sp = int(src['sp'] * multiplier)
    
    # Stock levels
    current_stock = random.choice([0, 0, 4, 7, 12, 18, 25, 36, 50, 75, 100, 150])
    sold_qty = random.randint(10, 250)
    status = "Out of Stock" if current_stock == 0 else ("Low Stock" if current_stock <= 10 else "In Stock")
    
    cat_prefix = src['category'][:3].upper().replace(" ", "")
    sku = f"MOD-{cat_prefix}-{i:04d}"
    mongo_id = hex(id_hex_base + i)[2:].zfill(24)
    
    final_480.append({
        "id": mongo_id,
        "slNo": i,
        "name": prod_name,
        "category": src['category'],
        "sku": sku,
        "unit": src['unit'],
        "mrp": mrp,
        "sp": sp,
        "currentStock": current_stock,
        "soldQty": sold_qty,
        "stockToAdd": 0,
        "status": status,
        "remarks": "Refill required" if current_stock <= 10 else "Normal stock"
    })

# 1. Save JSON
json_path = r'd:\modern\Modernstore-main\all_480_products_stock.json'
with open(json_path, 'w', encoding='utf-8') as f:
    json.dump(final_480, f, indent=2)

# 2. Build openpyxl Excel with all 480 Products
wb = openpyxl.Workbook()
ws1 = wb.active
ws1.title = "Stock Update Sheet"

header_fill = PatternFill(start_color="1B5E20", end_color="1B5E20", fill_type="solid")
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
add_col_fill = PatternFill(start_color="E65100", end_color="E65100", fill_type="solid")
accent_fill = PatternFill(start_color="FFF9C4", end_color="FFF9C4", fill_type="solid")
accent_font = Font(name="Calibri", size=11, bold=True, color="B71C1C")
thin_border = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9")
)

headers = [
    "Product ID (Do Not Change)",
    "Sl No",
    "Product Name",
    "Category",
    "SKU Code",
    "Unit",
    "MRP (Rs)",
    "Selling Price (Rs)",
    "Current Stock",
    "Sold Qty",
    "Stock To Add (Enter Qty Here)",
    "Projected Total Stock",
    "Stock Status",
    "Remarks"
]

ws1.append(headers)
for col_idx, cell in enumerate(ws1[1], 1):
    cell.fill = add_col_fill if col_idx == 11 else header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

for p in final_480:
    row_num = p['slNo'] + 1
    row_data = [
        p['id'],
        p['slNo'],
        p['name'],
        p['category'],
        p['sku'],
        p['unit'],
        p['mrp'],
        p['sp'],
        p['currentStock'],
        p['soldQty'],
        0, # Stock to add default
        f"=I{row_num}+K{row_num}", # Projected Total = Current Stock (col I) + Stock To Add (col K)
        p['status'],
        p['remarks']
    ]
    ws1.append(row_data)
    
    for c_idx, cell in enumerate(ws1[row_num], 1):
        cell.border = thin_border
        cell.font = Font(name="Calibri", size=10)
        cell.alignment = Alignment(vertical="center")
        
        if c_idx in (1, 2, 5, 6, 13):
            cell.alignment = Alignment(horizontal="center", vertical="center")
        elif c_idx in (7, 8, 9, 10, 11, 12):
            cell.alignment = Alignment(horizontal="right", vertical="center")
            cell.number_format = "#,##0"
            
        if c_idx == 11: # Stock to add
            cell.fill = accent_fill
            cell.font = accent_font

# Widths
ws1.column_dimensions['A'].width = 28
ws1.column_dimensions['B'].width = 8
ws1.column_dimensions['C'].width = 34
ws1.column_dimensions['D'].width = 22
ws1.column_dimensions['E'].width = 16
ws1.column_dimensions['F'].width = 10
ws1.column_dimensions['G'].width = 12
ws1.column_dimensions['H'].width = 15
ws1.column_dimensions['I'].width = 14
ws1.column_dimensions['J'].width = 12
ws1.column_dimensions['K'].width = 26
ws1.column_dimensions['L'].width = 20
ws1.column_dimensions['M'].width = 14
ws1.column_dimensions['N'].width = 18

# Instructions Sheet
ws2 = wb.create_sheet(title="Instructions & Field Guide")
ws2.views.sheetView[0].showGridLines = True
ws2.append(["MODERN STORE - 480 PRODUCTS BULK STOCK UPDATE INSTRUCTIONS"])
ws2.append([])
ws2.append(["Column Header", "Description", "Required / Action", "Example"])
ws2.append(["Product ID", "Unique database key used by backend. Do NOT alter.", "LOCKED", "67ec29000000000000000001"])
ws2.append(["Product Name", "Official registered product title.", "Read Only", "Fresh Shimla Apple"])
ws2.append(["Category", "Item category.", "Read Only", "Fresh Fruits"])
ws2.append(["Current Stock", "Live quantity currently remaining in the store.", "Read Only", "25"])
ws2.append(["Stock To Add", "ENTER THE NEW INCOMING STOCK QUANTITY HERE.", "TYPE HERE (Numbers)", "Enter 50 to add 50 units"])
ws2.append(["Projected Total", "Auto formula summing Current Stock + Stock To Add.", "Formula", "=I2+K2"])

ws2['A1'].font = Font(name="Calibri", size=14, bold=True, color="1B5E20")
for c in ws2[3]:
    c.font = header_font
    c.fill = header_fill
    c.alignment = Alignment(horizontal="center", vertical="center")

for col in ws2.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws2.column_dimensions[col_letter].width = max(max_len + 3, 14)

excel_path = r'd:\modern\Modernstore-main\ModernStore_All_480_Products_Stock.xlsx'
wb.save(excel_path)
print("Successfully generated 480 products Excel:", excel_path)
